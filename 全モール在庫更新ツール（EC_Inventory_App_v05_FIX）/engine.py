from __future__ import annotations
import csv, math, re, shutil
import json
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple, Optional
import openpyxl

ENCODINGS = ("cp932", "utf-8-sig", "shift_jis", "utf-8")

@dataclass
class ParseResult:
    raw_sku: str
    webcd: Optional[str]
    operation: str
    factor: int = 1
    note: str = ""

@dataclass
class AuditRecord:
    channel: str
    sku: str
    webcd: str
    wholesale_stock: int
    t2k_stock: int
    available_stock: int
    old_stock: int
    output_stock: int
    operation: str
    factor: int
    status: str
    note: str = ""

def _norm(v) -> str:
    return str(v or "").strip().replace("　"," ").strip()

def to_int(v, default=0):
    try:
        if v is None or str(v).strip()=="":
            return default
        return math.floor(float(str(v).replace(",","")))
    except Exception:
        return default

def read_csv(path: str|Path):
    path = Path(path)
    err=None
    for enc in ENCODINGS:
        try:
            with path.open("r",encoding=enc,newline="") as f:
                sample=f.read(4096); f.seek(0)
                delim="\t" if ("\t" in sample.splitlines()[0]) else ","
                r=csv.DictReader(f,delimiter=delim)
                rows=list(r)
                return list(r.fieldnames or []),rows,enc,delim
        except UnicodeDecodeError as e:
            err=e
    raise ValueError(f"文字コードを判定できません: {path.name}: {err}")

def write_csv(path,headers,rows,encoding="cp932"):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding=encoding,newline="",errors="replace") as f:
        w=csv.DictWriter(f,fieldnames=headers,extrasaction="ignore")
        w.writeheader(); w.writerows(rows)

def parse_sku(raw: str) -> ParseResult:
    s=_norm(raw).replace(" ","")
    if not s:
        return ParseResult(raw,None,"manual",1,"空SKU")
    if re.fullmatch(r"\d{7}",s):
        return ParseResult(raw,s,"quote",1,"7桁数字＝見積りコード")
    if re.fullmatch(r"\d{6}",s):
        return ParseResult(raw,s,"direct",1,"6桁WebCD")
    m=re.fullmatch(r"(\d{6})-(\d+)",s)
    if m:
        return ParseResult(raw,m.group(1),"direct",1,f"枝番 -{m.group(2)} を無視")
    m=re.fullmatch(r"(\d{6})[xXｘ×](\d+)",s)
    if m:
        n=int(m.group(2))
        return ParseResult(raw,m.group(1),"divide" if n>0 else "manual",n,"xN＝在庫÷N")
    m=re.fullmatch(r"(\d{6})(?:m[xXｘ×]|[xXｘ×]m)(\d+)",s,re.I)
    if m:
        n=int(m.group(2))
        return ParseResult(raw,m.group(1),"divide" if n>0 else "manual",n,"mxN/xmN＝在庫÷N")
    m=re.fullmatch(r"(\d{6})(?:/|[dD])(\d+)",s)
    if m:
        n=int(m.group(2))
        return ParseResult(raw,m.group(1),"multiply" if n>0 else "manual",n,"dN・/N＝在庫×N")
    return ParseResult(raw,None,"manual",1,"未定義SKU形式")

def latest_date_sheet(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_vba=True)
    c=[]
    for n in wb.sheetnames:
        if re.fullmatch(r"\d{8}",str(n)):
            try: c.append((datetime.strptime(n,"%Y%m%d"),n))
            except ValueError: pass
    if not c:
        raise ValueError("卸在庫ファイルにYYYYMMDD形式のシートがありません")
    return max(c)[1]

def load_wholesale(path):
    sheet=latest_date_sheet(path)
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_vba=True)
    ws=wb[sheet]
    out={}; warnings=[]
    # B = WebCD, M = 調整在庫
    for row in ws.iter_rows(min_row=2,min_col=1,max_col=13,values_only=True):
        web=_norm(row[1] if len(row)>1 else "")
        if not web: continue
        try: web=str(int(float(web)))
        except Exception: pass
        stock=to_int(row[12] if len(row)>12 else 0,0)
        if web in out: warnings.append(f"卸WebCD重複:{web}")
        out[web]=stock
    return out,sheet,warnings

def load_logiless(path):
    h,rows,_,_=read_csv(path)
    for req in ("商品コード","フリー在庫"):
        if req not in h: raise ValueError(f"LOGILESS必須列不足: {req}")
    out={}; warnings=[]
    for r in rows:
        web=_norm(r.get("商品コード"))
        if not web: continue
        if web in out: warnings.append(f"LOGILESS商品コード重複:{web}")
        out[web]=to_int(r.get("フリー在庫"),0)
    return out,warnings

def build_master(wholesale,t2k):
    keys=set(wholesale)|set(t2k)
    return {k:{
        "wholesale":wholesale.get(k,0),
        "t2k":t2k.get(k,0),
        "available":max(0,wholesale.get(k,0)+t2k.get(k,0))
    } for k in keys}

def calc(parsed,master,quote_inventory=None):
    empty={"wholesale":0,"t2k":0,"available":0}
    quote_inventory=quote_inventory or {}
    if parsed.operation=="quote":
        q=quote_inventory.get(parsed.raw_sku)
        if q is None:
            return None,"見積り構成未登録",empty
        info={"wholesale":0,"t2k":0,"available":int(q["available"])}
        return int(q["available"]),"OK",info
    if parsed.operation=="manual" or not parsed.webcd:
        return None,parsed.note,empty
    if parsed.webcd not in master:
        return None,"WebCD未一致",empty
    info=master[parsed.webcd]; a=info["available"]
    if parsed.operation=="direct": new=a
    elif parsed.operation=="divide":
        if parsed.factor<=0: return None,"除数不正",info
        new=math.floor(a/parsed.factor)
    elif parsed.operation=="multiply":
        if parsed.factor<=0: return None,"乗数不正",info
        new=a*parsed.factor
    else: return None,"未定義演算",info
    return max(0,int(new)),"OK",info

def audit_record(channel,sku,p,info,old,new,status):
    return AuditRecord(channel,sku,p.webcd or "",info["wholesale"],info["t2k"],
                       info["available"],old,new if new is not None else old,
                       p.operation,p.factor,status,p.note)

def update_mercari(src,out,master,audit,quote_inventory=None):
    h,rows,enc,_=read_csv(src)
    for r in rows:
        for i in range(1,11):
            c=f"SKU{i}_商品管理コード"; q=f"SKU{i}_現在の在庫数"
            if c not in h or q not in h: continue
            sku=_norm(r.get(c))
            if not sku: continue
            old=to_int(r.get(q),0); p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
            if new is not None: r[q]=str(new)
            audit.append(audit_record("Mercari",sku,p,info,old,new,status))
    write_csv(out,h,rows,enc if enc in ("cp932","utf-8-sig") else "cp932")

def update_amazon(report,template,out_xlsm,out_txt,master,audit,quote_inventory=None):
    h,rows,_,_=read_csv(report)
    sku_col="出品者SKU" if "出品者SKU" in h else h[0]
    qty_col="数量" if "数量" in h else (h[1] if len(h)>1 else None)
    computed=[]
    for r in rows:
        sku=_norm(r.get(sku_col))
        if not sku: continue
        old=to_int(r.get(qty_col),0) if qty_col else 0
        p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
        outqty=new if new is not None else old
        computed.append((sku,outqty))
        audit.append(audit_record("Amazon",sku,p,info,old,new,status))
    shutil.copy2(template,out_xlsm)
    wb=openpyxl.load_workbook(out_xlsm,keep_vba=True)
    if "テンプレート" not in wb.sheetnames:
        raise ValueError("Amazonテンプレートに「テンプレート」シートがありません")
    ws=wb["テンプレート"]
    for rr in range(7,max(ws.max_row,7)+1):
        for cc in (1,2,3):
            ws.cell(rr,cc).value=None
    for rr,(sku,qty) in enumerate(computed,start=7):
        ws.cell(rr,1).value=sku
        ws.cell(rr,2).value="DEFAULT"
        ws.cell(rr,3).value=int(qty)
    wb.save(out_xlsm)

    wb2=openpyxl.load_workbook(out_xlsm,read_only=True,data_only=False,keep_vba=True)
    ws2=wb2["テンプレート"]
    last=6+len(computed)
    with Path(out_txt).open("w",encoding="cp932",newline="",errors="replace") as f:
        w=csv.writer(f,delimiter="\t",lineterminator="\r\n")
        for row in ws2.iter_rows(min_row=1,max_row=last,min_col=1,max_col=ws2.max_column,values_only=True):
            w.writerow(["" if v is None else v for v in row])

def load_itemsub_map(path):
    h,rows,_,_=read_csv(path)
    main_col="メインデータの商品コード（楽天URL）"
    sys_col="システム連携用SKU番号"
    if main_col not in h or sys_col not in h:
        raise ValueError("ir-itemsubの必須列（メインデータの商品コード（楽天URL）/システム連携用SKU番号）がありません")
    out={}
    dup=[]
    for r in rows:
        main=_norm(r.get(main_col))
        sys=_norm(r.get(sys_col))
        if not main or not sys:
            continue
        if main in out and out[main] != sys:
            dup.append(main)
        else:
            out[main]=sys
    return out,sorted(set(dup))

def update_itemrobot(ir_item,ir_sel,ir_itemsub,out_item,out_sel,master,audit,quote_inventory=None):
    # 選択項目あり: 楽天 システム連携用SKU番号をWebCD照合キーに使用
    h,rows,enc,_=read_csv(ir_sel)
    sys_col="楽天　システム連携用SKU番号"
    stock_col="項目選択肢別在庫用在庫数"
    if sys_col not in h: raise ValueError(f"ir-selectionに「{sys_col}」がありません")
    if stock_col not in h: raise ValueError(f"ir-selectionに「{stock_col}」がありません")
    for r in rows:
        sku=_norm(r.get(sys_col))
        if not sku: continue
        old=to_int(r.get(stock_col),0); p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
        if new is not None: r[stock_col]=str(new)
        audit.append(audit_record("itemRobot-selection",sku,p,info,old,new,status))
    write_csv(out_sel,h,rows,enc if enc in ("cp932","utf-8-sig") else "cp932")

    # 選択項目なし:
    # ir-itemの商品コード（楽天URL） → ir-itemsubのメインデータの商品コード（楽天URL）
    # → itemsubのシステム連携用SKU番号 → WebCD
    itemsub_map,itemsub_dups=load_itemsub_map(ir_itemsub)

    h2,rows2,enc2,_=read_csv(ir_item)
    code_col="商品コード（楽天URL）"; stock_col2="在庫数"
    if code_col not in h2 or stock_col2 not in h2:
        raise ValueError("ir-itemの必須列（商品コード（楽天URL）/在庫数）がありません")

    for r in rows2:
        item_code=_norm(r.get(code_col))
        if not item_code: continue
        old=to_int(r.get(stock_col2),0)

        if item_code in itemsub_dups:
            p=ParseResult(item_code,None,"manual",1,"itemsubで同一メイン商品コードに複数システムSKU")
            info={"wholesale":0,"t2k":0,"available":0}
            new=None; status="itemsub重複・照合保留"
            audit.append(audit_record("itemRobot-item",item_code,p,info,old,new,status))
            continue

        sys_sku=itemsub_map.get(item_code,"")
        if sys_sku:
            p=parse_sku(sys_sku); new,status,info=calc(p,master,quote_inventory)
            if new is not None:
                r[stock_col2]=str(new)
            rec=audit_record("itemRobot-item",sys_sku,p,info,old,new,status)
            rec.note=(rec.note + f" / itemコード:{item_code} / itemsub経由").strip(" /")
            audit.append(rec)
            continue

        # itemsubに存在しない場合のみ、安全なフォールバック:
        # 商品コード自体が6桁かつWebCDとしてマスター一致する場合は直接更新
        p=parse_sku(item_code); new,status,info=calc(p,master,quote_inventory)
        safe=bool(re.fullmatch(r"\d{6}",item_code) and p.webcd in master)
        if safe and new is not None:
            r[stock_col2]=str(new)
            rec=audit_record("itemRobot-item",item_code,p,info,old,new,"OK")
            rec.note=(rec.note+" / itemsubなし・商品コード直接照合").strip(" /")
            audit.append(rec)
        else:
            new=None
            if status=="OK": status="itemsubなし・照合保留"
            elif status=="WebCD未一致": status="itemsubなし・WebCD未一致"
            rec=audit_record("itemRobot-item",item_code,p,info,old,new,status)
            rec.note=(rec.note+" / itemsubにメイン商品コードなし").strip(" /")
            audit.append(rec)

    write_csv(out_item,h2,rows2,enc2 if enc2 in ("cp932","utf-8-sig") else "cp932")

def save_audit(path,audit):
    headers=["channel","sku","webcd","wholesale_stock","t2k_stock","available_stock","old_stock",
             "output_stock","operation","factor","status","note"]
    write_csv(path,headers,[asdict(x) for x in audit],"utf-8-sig")

def save_master(path,master):
    headers=["WebCD","卸調整在庫","T2Kフリー在庫","販売可能在庫"]
    rows=[{"WebCD":k,"卸調整在庫":v["wholesale"],"T2Kフリー在庫":v["t2k"],"販売可能在庫":v["available"]}
          for k,v in sorted(master.items())]
    write_csv(path,headers,rows,"utf-8-sig")

def summarize(audit):
    from collections import Counter
    statuses=Counter(r.status for r in audit)
    channels=Counter(r.channel for r in audit)
    ok=sum(1 for r in audit if r.status=="OK")
    changed=sum(1 for r in audit if r.status=="OK" and r.old_stock!=r.output_stock)
    return {"total":len(audit),"ok":ok,"changed":changed,"statuses":dict(statuses),"channels":dict(channels)}

def validate_inputs(inputs, channels=None):
    """Validate file existence and required columns before any stock calculation."""
    channels=set(channels or ["itemrobot","amazon","mercari"])
    errors=[]
    warnings=[]

    required=["wholesale","logiless"]
    if "itemrobot" in channels:
        required += ["ir_item","ir_selection","ir_itemsub"]
    if "amazon" in channels:
        required += ["amazon_report","amazon_template"]
    if "mercari" in channels:
        required += ["mercari"]

    for k in required:
        p=Path(inputs.get(k,""))
        if not p.exists():
            errors.append(f"{k}: ファイルが見つかりません")

    if errors:
        return errors,warnings

    # Wholesale
    try:
        s=latest_date_sheet(inputs["wholesale"])
        warnings.append(f"卸在庫使用予定シート: {s}")
    except Exception as e:
        errors.append(f"卸在庫: {e}")

    # LOGILESS
    try:
        h,_,_,_=read_csv(inputs["logiless"])
        for c in ("商品コード","フリー在庫"):
            if c not in h:
                errors.append(f"LOGILESS: 必須列「{c}」がありません")
    except Exception as e:
        errors.append(f"LOGILESS: {e}")

    if "itemrobot" in channels:
        try:
            h,_,_,_=read_csv(inputs["ir_item"])
            for c in ("商品コード（楽天URL）","在庫数"):
                if c not in h:
                    errors.append(f"ir-item: 必須列「{c}」がありません")
        except Exception as e:
            errors.append(f"ir-item: {e}")
        try:
            h,_,_,_=read_csv(inputs["ir_selection"])
            for c in ("楽天　システム連携用SKU番号","項目選択肢別在庫用在庫数"):
                if c not in h:
                    errors.append(f"ir-selection: 必須列「{c}」がありません")
        except Exception as e:
            errors.append(f"ir-selection: {e}")
        try:
            h,_,_,_=read_csv(inputs["ir_itemsub"])
            for c in ("メインデータの商品コード（楽天URL）","システム連携用SKU番号"):
                if c not in h:
                    errors.append(f"ir-itemsub: 必須列「{c}」がありません")
        except Exception as e:
            errors.append(f"ir-itemsub: {e}")

    if "amazon" in channels:
        try:
            h,_,_,_=read_csv(inputs["amazon_report"])
            if "出品者SKU" not in h:
                errors.append("Amazon出品レポート: 「出品者SKU」列がありません")
        except Exception as e:
            errors.append(f"Amazon出品レポート: {e}")
        try:
            wb=openpyxl.load_workbook(inputs["amazon_template"],read_only=True,keep_vba=True)
            if "テンプレート" not in wb.sheetnames:
                errors.append("Amazonテンプレート: 「テンプレート」シートがありません")
        except Exception as e:
            errors.append(f"Amazonテンプレート: {e}")

    if "mercari" in channels:
        try:
            h,_,_,_=read_csv(inputs["mercari"])
            if "SKU1_商品管理コード" not in h:
                errors.append("メルカリ: 「SKU1_商品管理コード」列がありません")
        except Exception as e:
            errors.append(f"メルカリ: {e}")

    return errors,warnings

def analyze_audit(audit, severe_drop_pct=80):
    """Produce channel/status summaries plus large stock decrease warnings."""
    status_by_channel=defaultdict(Counter)
    change_by_channel=Counter()
    severe=[]
    for r in audit:
        status_by_channel[r.channel][r.status]+=1
        if r.status=="OK" and r.old_stock != r.output_stock:
            change_by_channel[r.channel]+=1
        if r.status=="OK" and r.old_stock > 0 and r.output_stock < r.old_stock:
            drop=(r.old_stock-r.output_stock)/r.old_stock*100
            if drop >= severe_drop_pct:
                severe.append({
                    "channel":r.channel,"sku":r.sku,"webcd":r.webcd,
                    "old_stock":r.old_stock,"new_stock":r.output_stock,
                    "drop_pct":round(drop,1)
                })
    return {
        "status_by_channel":{k:dict(v) for k,v in status_by_channel.items()},
        "change_by_channel":dict(change_by_channel),
        "severe_drops":severe,
    }

def save_precheck_report(path, audit):
    headers=["channel","sku","webcd","old_stock","new_stock","status","note"]
    rows=[]
    for r in audit:
        if r.status!="OK":
            rows.append({
                "channel":r.channel,"sku":r.sku,"webcd":r.webcd,
                "old_stock":r.old_stock,"new_stock":r.output_stock,
                "status":r.status,"note":r.note
            })
    write_csv(path,headers,rows,"utf-8-sig")

def load_quote_master(path):
    """Load optional 7-digit estimate/BOM master. CSV or XLSX.
    Required: 見積り番号, Webコード, 必要数量
    Optional: 見積り商品名, 有効, 備考
    """
    if not path or not Path(path).exists():
        return {},[]
    p=Path(path)
    rows=[]
    if p.suffix.lower() in (".xlsx",".xlsm"):
        wb=openpyxl.load_workbook(p,read_only=True,data_only=True,keep_vba=p.suffix.lower()==".xlsm")
        ws=wb[wb.sheetnames[0]]
        vals=list(ws.iter_rows(values_only=True))
        if not vals:
            return {},["見積り構成マスターが空です"]
        headers=[_norm(x) for x in vals[0]]
        for row in vals[1:]:
            rows.append({headers[i]: row[i] if i<len(row) else "" for i in range(len(headers))})
    else:
        headers,rows,_,_=read_csv(p)

    required=("見積り番号","Webコード","必要数量")
    for c in required:
        if c not in (headers if p.suffix.lower() not in (".xlsx",".xlsm") else [_norm(x) for x in vals[0]]):
            raise ValueError(f"見積り構成マスター必須列不足: {c}")

    out=defaultdict(list); warnings=[]
    for r in rows:
        quote=_norm(r.get("見積り番号"))
        web=_norm(r.get("Webコード"))
        qty=to_int(r.get("必要数量"),0)
        active=_norm(r.get("有効","1")).lower()
        if active in ("0","false","無効","off","no"):
            continue
        if not re.fullmatch(r"\d{7}",quote):
            if quote: warnings.append(f"見積り番号形式不正:{quote}")
            continue
        if not re.fullmatch(r"\d{6}",web):
            warnings.append(f"見積り{quote}: Webコード形式不正:{web}")
            continue
        if qty<=0:
            warnings.append(f"見積り{quote}: 必要数量不正:{web}={qty}")
            continue
        out[quote].append({
            "webcd":web,
            "qty":qty,
            "name":_norm(r.get("見積り商品名","")),
            "note":_norm(r.get("備考",""))
        })
    return dict(out),warnings

def build_quote_inventory(master, quote_master):
    result={}
    warnings=[]
    for quote, comps in quote_master.items():
        if not comps:
            continue
        possible=[]
        missing=[]
        limiting=[]
        for c in comps:
            web=c["webcd"]; qty=c["qty"]
            if web not in master:
                missing.append(web)
                continue
            av=master[web]["available"]
            units=math.floor(av/qty)
            possible.append((units,web,av,qty))
        if missing:
            result[quote]={
                "available":0,"reason":"構成WebCD未一致: "+",".join(missing),
                "components":comps,"name":comps[0].get("name","")
            }
            warnings.append(f"見積り{quote}: 構成WebCD未一致 {','.join(missing)}")
            continue
        if not possible:
            result[quote]={"available":0,"reason":"構成なし","components":comps,"name":comps[0].get("name","")}
            continue
        min_units=min(x[0] for x in possible)
        limiting=[x for x in possible if x[0]==min_units]
        reason=" / ".join([f"{x[1]}:在庫{x[2]}÷必要{x[3]}={x[0]}" for x in limiting])
        result[quote]={
            "available":int(min_units),
            "reason":reason,
            "components":comps,
            "name":comps[0].get("name","")
        }
    return result,warnings

def load_wholesale_names(path):
    sheet=latest_date_sheet(path)
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_vba=True)
    ws=wb[sheet]
    names={}
    # B=WebCD, C=商品名想定
    for row in ws.iter_rows(min_row=2,min_col=1,max_col=3,values_only=True):
        web=_norm(row[1] if len(row)>1 else "")
        if not web: continue
        try: web=str(int(float(web)))
        except Exception: pass
        name=_norm(row[2] if len(row)>2 else "")
        if name and web not in names:
            names[web]=name
    return names

def _state_path(base_output):
    p=Path(base_output)/"_system"
    p.mkdir(parents=True,exist_ok=True)
    return p/"inventory_state.json"

def load_inventory_state(base_output):
    p=_state_path(base_output)
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}

def save_inventory_state(base_output,state):
    p=_state_path(base_output)
    p.write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding="utf-8")

def shortage_status(prev,current):
    if prev is None:
        return "初回欠品" if current==0 else "正常"
    prev=int(prev); current=int(current)
    if prev>0 and current==0: return "新規欠品"
    if prev==0 and current==0: return "継続欠品"
    if prev==0 and current>0: return "欠品解消"
    return "正常"

def active_product_codes(audit, quote_inventory):
    web=set(); quotes=set()
    for r in audit:
        if re.fullmatch(r"\d{6}",r.webcd or ""):
            web.add(r.webcd)
        if re.fullmatch(r"\d{7}",r.sku or "") and r.sku in quote_inventory:
            quotes.add(r.sku)
    return web,quotes

def build_product_shortage_rows(master, quote_inventory, audit, previous_state, names):
    web_codes,quote_codes=active_product_codes(audit,quote_inventory)
    rows=[]; new_state={}
    for web in sorted(web_codes):
        cur=int(master.get(web,{}).get("available",0))
        key=f"WEB:{web}"
        prev=previous_state.get(key,{}).get("stock")
        st=shortage_status(prev,cur)
        new_state[key]={"stock":cur,"type":"Web6桁","code":web}
        if st!="正常":
            rows.append({
                "種別":"Web6桁","商品コード":web,"商品名":names.get(web,""),
                "前回在庫":"" if prev is None else int(prev),"今回在庫":cur,
                "欠品区分":st,"欠品原因":"卸調整在庫＋T2Kフリー在庫が0" if cur==0 else "在庫復活"
            })
    for q in sorted(quote_codes):
        cur=int(quote_inventory[q]["available"])
        key=f"QUOTE:{q}"
        prev=previous_state.get(key,{}).get("stock")
        st=shortage_status(prev,cur)
        new_state[key]={"stock":cur,"type":"見積り7桁","code":q}
        if st!="正常":
            rows.append({
                "種別":"見積り7桁","商品コード":q,"商品名":quote_inventory[q].get("name",""),
                "前回在庫":"" if prev is None else int(prev),"今回在庫":cur,
                "欠品区分":st,"欠品原因":quote_inventory[q].get("reason","") if cur==0 else "構成商品の在庫復活"
            })
    order={"新規欠品":0,"欠品解消":1,"継続欠品":2,"初回欠品":3}
    rows.sort(key=lambda r:(order.get(r["欠品区分"],9),r["種別"],r["商品コード"]))
    return rows,new_state

def build_sku_shortage_rows(audit, previous_state):
    rows=[]; new_state={}
    for r in audit:
        if r.status!="OK":
            continue
        key=f"SKU:{r.channel}:{r.sku}"
        cur=int(r.output_stock)
        prev=previous_state.get(key,{}).get("stock")
        st=shortage_status(prev,cur)
        new_state[key]={"stock":cur,"type":"販売SKU","code":r.sku,"channel":r.channel}
        if st!="正常":
            rows.append({
                "モール":r.channel,"販売SKU":r.sku,"Webコード/見積り":r.webcd,
                "前回在庫":"" if prev is None else int(prev),"今回在庫":cur,
                "欠品区分":st,"計算":r.operation,"係数":r.factor,"備考":r.note
            })
    order={"新規欠品":0,"欠品解消":1,"継続欠品":2,"初回欠品":3}
    rows.sort(key=lambda x:(order.get(x["欠品区分"],9),x["モール"],x["販売SKU"]))
    return rows,new_state

def save_shortage_csvs(outdir,date,product_rows,sku_rows):
    p1=Path(outdir)/f"欠品商品_{date}.csv"
    p2=Path(outdir)/f"販売SKU欠品_{date}.csv"
    h1=["種別","商品コード","商品名","前回在庫","今回在庫","欠品区分","欠品原因"]
    h2=["モール","販売SKU","Webコード/見積り","前回在庫","今回在庫","欠品区分","計算","係数","備考"]
    write_csv(p1,h1,product_rows,"utf-8-sig")
    write_csv(p2,h2,sku_rows,"utf-8-sig")
    return str(p1),str(p2)

def precheck(inputs,channels=None):
    channels=set(channels or ["itemrobot","amazon","mercari"])
    wholesale,sheet,w1=load_wholesale(inputs["wholesale"])
    t2k,w2=load_logiless(inputs["logiless"])
    master=build_master(wholesale,t2k)
    quote_master,qw=load_quote_master(inputs.get("quote_master",""))
    quote_inventory,qiw=build_quote_inventory(master,quote_master)
    # lightweight audit uses temporary in-memory logic without writing outputs
    audit=[]
    if "mercari" in channels:
        h,rows,_,_=read_csv(inputs["mercari"])
        for r in rows:
            for i in range(1,11):
                c=f"SKU{i}_商品管理コード"; q=f"SKU{i}_現在の在庫数"
                if c not in h: continue
                sku=_norm(r.get(c))
                if not sku: continue
                old=to_int(r.get(q),0); p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
                audit.append(audit_record("Mercari",sku,p,info,old,new,status))
    if "amazon" in channels:
        h,rows,_,_=read_csv(inputs["amazon_report"])
        sc="出品者SKU" if "出品者SKU" in h else h[0]; qc="数量" if "数量" in h else None
        for r in rows:
            sku=_norm(r.get(sc))
            if not sku: continue
            old=to_int(r.get(qc),0) if qc else 0
            p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
            audit.append(audit_record("Amazon",sku,p,info,old,new,status))
    if "itemrobot" in channels:
        h,rows,_,_=read_csv(inputs["ir_selection"])
        sc="楽天　システム連携用SKU番号"; qc="項目選択肢別在庫用在庫数"
        for r in rows:
            sku=_norm(r.get(sc))
            if not sku: continue
            old=to_int(r.get(qc),0); p=parse_sku(sku); new,status,info=calc(p,master,quote_inventory)
            audit.append(audit_record("itemRobot-selection",sku,p,info,old,new,status))

        itemsub_map,itemsub_dups=load_itemsub_map(inputs["ir_itemsub"])
        h2,rows2,_,_=read_csv(inputs["ir_item"])
        for r in rows2:
            item_code=_norm(r.get("商品コード（楽天URL）"))
            if not item_code: continue
            old=to_int(r.get("在庫数"),0)

            if item_code in itemsub_dups:
                p=ParseResult(item_code,None,"manual",1,"itemsubで同一メイン商品コードに複数システムSKU")
                info={"wholesale":0,"t2k":0,"available":0}
                audit.append(audit_record("itemRobot-item",item_code,p,info,old,None,"itemsub重複・照合保留"))
                continue

            sys_sku=itemsub_map.get(item_code,"")
            if sys_sku:
                p=parse_sku(sys_sku); new,status,info=calc(p,master,quote_inventory)
                rec=audit_record("itemRobot-item",sys_sku,p,info,old,new,status)
                rec.note=(rec.note+f" / itemコード:{item_code} / itemsub経由").strip(" /")
                audit.append(rec)
            else:
                p=parse_sku(item_code); new,status,info=calc(p,master,quote_inventory)
                safe=bool(re.fullmatch(r"\d{6}",item_code) and p.webcd in master)
                if not safe:
                    new=None
                    if status=="OK": status="itemsubなし・照合保留"
                    elif status=="WebCD未一致": status="itemsubなし・WebCD未一致"
                rec=audit_record("itemRobot-item",item_code,p,info,old,new,status if not safe else "OK")
                rec.note=(rec.note+" / itemsubにメイン商品コードなし").strip(" /")
                audit.append(rec)
    return {"sheet":sheet,"master_count":len(master),"quote_count":len(quote_inventory),"warnings":w1+w2+qw+qiw,"summary":summarize(audit),"analysis":analyze_audit(audit),"audit":audit}

def run_all(inputs,output_dir,channels=None, create_run_folder=True):
    channels=set(channels or ["itemrobot","amazon","mercari"])
    base=Path(output_dir); base.mkdir(parents=True,exist_ok=True)
    if create_run_folder:
        stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
        outdir=base/f"EC在庫更新_{stamp}"
        outdir.mkdir(parents=True,exist_ok=True)
    else:
        outdir=base
    wholesale,sheet,w1=load_wholesale(inputs["wholesale"])
    t2k,w2=load_logiless(inputs["logiless"])
    master=build_master(wholesale,t2k)
    quote_master,qw=load_quote_master(inputs.get("quote_master",""))
    quote_inventory,qiw=build_quote_inventory(master,quote_master)
    audit=[]
    date=datetime.now().strftime("%Y%m%d")
    save_master(outdir/f"統合在庫マスター_{date}.csv",master)

    outputs={"wholesale_sheet":sheet,"warnings":w1+w2}
    if "itemrobot" in channels:
        a=outdir/f"ir-item_在庫更新_{date}.csv"; b=outdir/f"ir-selection_在庫更新_{date}.csv"
        update_itemrobot(inputs["ir_item"],inputs["ir_selection"],inputs["ir_itemsub"],a,b,master,audit,quote_inventory)
        outputs["ir_item"]=str(a); outputs["ir_selection"]=str(b)
    if "amazon" in channels:
        x=outdir/f"Amazon_PriceAndQuantity_{date}.xlsm"
        t=outdir/f"Amazon在庫更新_{date}.txt"
        update_amazon(inputs["amazon_report"],inputs["amazon_template"],x,t,master,audit,quote_inventory)
        outputs["amazon_xlsm"]=str(x); outputs["amazon_txt"]=str(t)
    if "mercari" in channels:
        m=outdir/f"メルカリ在庫更新_{date}.csv"
        update_mercari(inputs["mercari"],m,master,audit,quote_inventory); outputs["mercari"]=str(m)

    audit_path=outdir/f"在庫更新監査_{date}.csv"
    save_audit(audit_path,audit); outputs["audit"]=str(audit_path)
    outputs["summary"]=summarize(audit)
    outputs["analysis"]=analyze_audit(audit)
    precheck_path=outdir/f"要確認SKU_{date}.csv"
    save_precheck_report(precheck_path,audit)
    outputs["precheck_report"]=str(precheck_path)

    previous_state=load_inventory_state(base)
    names=load_wholesale_names(inputs["wholesale"])
    product_rows,product_state=build_product_shortage_rows(master,quote_inventory,audit,previous_state,names)
    sku_rows,sku_state=build_sku_shortage_rows(audit,previous_state)
    shortage_product,shortage_sku=save_shortage_csvs(outdir,date,product_rows,sku_rows)
    merged=dict(previous_state); merged.update(product_state); merged.update(sku_state)
    save_inventory_state(base,merged)

    outputs["shortage_product"]=shortage_product
    outputs["shortage_sku"]=shortage_sku
    outputs["shortage_counts"]={
        "新規欠品":sum(1 for r in product_rows if r["欠品区分"]=="新規欠品"),
        "継続欠品":sum(1 for r in product_rows if r["欠品区分"]=="継続欠品"),
        "欠品解消":sum(1 for r in product_rows if r["欠品区分"]=="欠品解消"),
        "初回欠品":sum(1 for r in product_rows if r["欠品区分"]=="初回欠品"),
    }
    outputs["quote_count"]=len(quote_inventory)
    outputs["warnings"]=outputs.get("warnings",[])+qw+qiw
    outputs["output_dir"]=str(outdir)
    return outputs
